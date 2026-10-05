
import os
import re
import glob
import streamlit as st
from pptx import Presentation
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

st.set_page_config(
    page_title="AI Engineering Learning Assistant",
    page_icon="🎓",
    layout="centered"
)

st.title("🎓 RJ's Mobile Communication Teaching Assistant")
st.caption("AI-powered learning assistant for Advanced Mobile Communication")

st.info("📚 Learn from faculty teaching PPTs • Ask questions • Get simple explanations • Generate notes, MCQs and 5-mark answers")

PPT_FOLDER = "PPTs"

@st.cache_data
def load_ppts():
    files = sorted(glob.glob(os.path.join(PPT_FOLDER, "*.pptx")))
    slides = []

    for file in files:
        unit = os.path.splitext(os.path.basename(file))[0]
        prs = Presentation(file)

        for slide_no, slide in enumerate(prs.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    t = shape.text.strip()
                    if t:
                        parts.append(t)

            text = "\n".join(parts).strip()
            if text:
                slides.append({
                    "unit": unit,
                    "file": os.path.basename(file),
                    "slide": slide_no,
                    "text": text
                })

    return slides

slides = load_ppts()

if not slides:
    st.error("No PPT files found. Put your .pptx files inside the PPTs folder.")
    st.stop()

documents = [s["text"] for s in slides]

@st.cache_resource
def build_search_engine(docs):
    vectorizer = TfidfVectorizer(stop_words="english")
    matrix = vectorizer.fit_transform(docs)
    return vectorizer, matrix

vectorizer, matrix = build_search_engine(documents)

units = sorted(list(dict.fromkeys(s["unit"] for s in slides)))

st.sidebar.header("📚 Learning Settings")

selected_unit = st.sidebar.selectbox(
    "Select Unit",
    units
)

mode = st.sidebar.radio(
    "Choose an activity",
    [
        "🤖 Ask Question",
        "💡 Explain Simply",
        "📝 Study Notes",
        "🎯 Generate MCQs",
        "✍️ 5-Mark Answer"
    ]
)

def clean(text):
    return re.sub(r"\s+", " ", text).strip()

def search(query, unit, n=5):
    indices = [
        i for i, s in enumerate(slides)
        if s["unit"] == unit
    ]

    if not indices:
        return []

    qv = vectorizer.transform([query])
    scores = cosine_similarity(qv, matrix[indices])[0]

    ranked = sorted(
        zip(indices, scores),
        key=lambda x: x[1],
        reverse=True
    )

    results = []

    for idx, score in ranked:
        if score <= 0:
            continue
        item = slides[idx].copy()
        item["score"] = float(score)
        results.append(item)
        if len(results) >= n:
            break

    return results

def source_text(results):
    if not results:
        return "No matching PPT slide was found."

    return "\n".join(
        f"- {r['file']} — Slide {r['slide']}"
        for r in results
    )

def answer_question(query, results):
    if not results:
        return "I could not find this topic in the selected unit."

    best = clean(results[0]["text"])

    if len(best) > 1200:
        best = best[:1200] + "..."

    return f"""
### 📚 Answer from the Teaching Material

{best}

### 📖 Source Slides
{source_text(results)}
"""

def explain_simply(query, results):
    if not results:
        return "I could not find this topic in the selected unit."

    text = clean(results[0]["text"])
    sentences = re.split(r"(?<=[.!?])\s+", text)
    sentences = [s for s in sentences if s.strip()]

    short = " ".join(sentences[:5])

    if len(short) > 900:
        short = short[:900] + "..."

    return f"""
### 💡 Simple Explanation

**Topic:** {query}

{short}

**In simple words:**  
The above explanation is taken directly from the selected teaching material and shortened for easier student understanding.

### 📖 Source Slides
{source_text(results)}
"""

def make_notes(query, results):
    if not results:
        return "No relevant notes were found."

    output = [f"### 📝 Study Notes: {query}", ""]

    for i, r in enumerate(results, start=1):
        text = clean(r["text"])
        if len(text) > 650:
            text = text[:650] + "..."

        output.append(f"**Point {i} — Slide {r['slide']}**")
        output.append(text)
        output.append("")

    output.append("### 📖 Sources")
    output.append(source_text(results))

    return "\n".join(output)

def make_mcqs(query, results):
    if not results:
        return "No relevant PPT content was found for MCQs."

    sentences = []

    for r in results:
        parts = re.split(r"(?<=[.!?])\s+", clean(r["text"]))
        for p in parts:
            if len(p.split()) >= 8:
                sentences.append(p)

    sentences = sentences[:5]

    if not sentences:
        return "Not enough structured PPT content to create MCQs."

    out = [f"### 🎯 Practice MCQs — {query}", ""]

    for i, sentence in enumerate(sentences, start=1):
        out.append(f"**Q{i}. Which statement is supported by the selected teaching material?**")
        out.append(f"A. {sentence[:220]}")
        out.append("B. This concept is unrelated to communication engineering.")
        out.append("C. This topic is used only in mechanical engineering.")
        out.append("D. None of the above.")
        out.append("**Correct answer: A**")
        out.append("")

    out.append("### 📖 Sources")
    out.append(source_text(results))

    return "\n".join(out)

def make_5mark(query, results):
    if not results:
        return "The topic was not found in the selected unit."

    text = " ".join(clean(r["text"]) for r in results)
    sentences = re.split(r"(?<=[.!?])\s+", text)
    sentences = [s for s in sentences if s.strip()]

    definition = sentences[0] if sentences else text[:500]
    explanation = " ".join(sentences[1:4])
    key_points = sentences[4:8]

    out = [
        f"### ✍️ 5-Mark Answer: {query}",
        "",
        "**1. Definition**",
        definition,
        "",
        "**2. Explanation**",
        explanation,
        "",
        "**3. Important Points**"
    ]

    for p in key_points:
        out.append(f"- {p}")

    out += [
        "",
        "**4. Conclusion**",
        "The concept should be understood with the relevant diagram, example, and terminology given in the teaching material.",
        "",
        "### 📖 Sources",
        source_text(results)
    ]

    return "\n".join(out)

st.write(f"**📚 Selected Unit:** `{selected_unit}`")

query = st.text_area(
    "❓ Enter your question or topic",
    placeholder="Example: What is Doppler Effect? Explain WSSUS. What is GSM?",
    height=110
)

if st.button("🚀 Ask Learning Assistant", type="primary", use_container_width=True):
    if not query.strip():
        st.warning("Please enter a question or topic.")
    else:
        with st.spinner("🔎 Searching the selected teaching material..."):
            results = search(query, selected_unit, 5)

        if mode == "🤖 Ask Question":
            response = answer_question(query, results)
        elif mode == "💡 Explain Simply":
            response = explain_simply(query, results)
        elif mode == "📝 Study Notes":
            response = make_notes(query, results)
        elif mode == "🎯 Generate MCQs":
            response = make_mcqs(query, results)
        else:
            response = make_5mark(query, results)

        st.markdown(response)

st.divider()

st.caption(
    f"📊 Knowledge base: {len(slides)} PPT slides across {len(units)} units."
)
